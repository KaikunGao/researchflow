"""研究报告生成（Word .docx）。

四种报告类型：
- company_deep_dive：单家公司财务拆解（核心数据表 + 点评段落 + 四张图）
- peer_comparison：多家公司在选定指标上的横向对比（表 + 图 + 排名）
- industry_overview：板块综述（最新报告期汇总表、子行业均值、涨跌幅前列、活跃告警）
- periodic_update：定期跟踪（每家公司最新报告期点评 + 核心表 + 告警）
"""

import io
import json
import sqlite3
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services import alerts as alert_service
from app.services import charts, repository
from app.services.catalog import METRIC_CATALOG, METRIC_MAP, MetricDef, flatten, format_value
from app.services.commentary import METHODOLOGY_NOTE, build_commentary
from app.services.industry import SUB_INDUSTRY_LABELS, SUB_INDUSTRY_ORDER, classify
from app.services.metrics import PeriodMetrics, build_series
from app.services.periods import Period
from app.services.settings_store import get_all as get_settings_all
from app.services.settings_store import source_line

REPORT_TYPE_LABELS = {
    "industry_overview": "行业综述",
    "company_deep_dive": "公司深度",
    "peer_comparison": "同业对比",
    "periodic_update": "定期跟踪",
}

DEFAULT_METRICS = ["revenue", "revenue_yoy", "net_profit", "net_profit_yoy", "gross_margin", "net_margin", "selling_ratio", "rd_ratio"]

Number = float | None


class ReportError(ValueError):
    pass


def _now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


# --------------------------------------------------------------------------- persistence


def list_reports(conn: sqlite3.Connection) -> list[dict]:
    rows = conn.execute("SELECT * FROM reports ORDER BY created_at DESC").fetchall()
    return [_row(r) for r in rows]


def get_report(conn: sqlite3.Connection, report_id: str) -> dict | None:
    row = conn.execute("SELECT * FROM reports WHERE id = ?", (report_id,)).fetchone()
    return _row(row) if row else None


def delete_report(conn: sqlite3.Connection, report_id: str) -> dict | None:
    report = get_report(conn, report_id)
    if report:
        conn.execute("DELETE FROM reports WHERE id = ?", (report_id,))
    return report


def _row(row: sqlite3.Row) -> dict:
    d = dict(row)
    d["codes"] = json.loads(d["codes"])
    d["metrics"] = json.loads(d["metrics"])
    d["report_type_label"] = REPORT_TYPE_LABELS.get(d["report_type"], d["report_type"])
    return d


# --------------------------------------------------------------------------- generation


def generate(conn: sqlite3.Connection, reports_dir, report_type: str, codes: list[str], period_end: str | None, metrics: list[str]) -> dict:
    if report_type not in REPORT_TYPE_LABELS:
        raise ReportError("未知的报告类型。")
    metrics = [m for m in metrics if m in METRIC_MAP] or DEFAULT_METRICS
    companies = {c["code"]: c for c in repository.list_companies(conn)}
    if report_type == "industry_overview" and not codes:
        codes = list(companies)
    codes = [c for c in codes if c in companies]
    if not codes:
        raise ReportError("请至少选择一家已导入数据的公司。")
    if report_type == "company_deep_dive" and len(codes) != 1:
        raise ReportError("公司深度报告一次只能选择一家公司。")
    if report_type == "peer_comparison" and len(codes) < 2:
        raise ReportError("同业对比至少需要两家公司。")

    settings = get_settings_all(conn)
    source = f"{source_line(settings)}，ResearchFlow 整理"
    series_map: dict[str, list[PeriodMetrics]] = {}
    for code in codes:
        facts = repository.get_company_facts(conn, code)
        if facts:
            series_map[code] = build_series(facts)
    if not series_map:
        raise ReportError("所选公司没有财务数据。")

    # 报告期：指定则取该期（没有的公司取其最新），否则取各公司最新
    def pick(series: list[PeriodMetrics]) -> PeriodMetrics:
        if period_end:
            for m in series:
                if m.period_end == period_end:
                    return m
        return series[-1]

    doc = Document()
    _setup_document(doc)
    builder = _Builder(doc, source, settings)

    if report_type == "company_deep_dive":
        code = codes[0]
        title = f"{companies[code]['name']}（{code}）{pick(series_map[code]).label}财务点评"
        builder.deep_dive(companies[code], series_map[code], pick(series_map[code]), metrics)
    elif report_type == "peer_comparison":
        names = "、".join(companies[c]["name"] for c in codes[:3]) + (f" 等 {len(codes)} 家" if len(codes) > 3 else "")
        title = f"{names}同业财务对比"
        builder.peer_comparison(companies, series_map, pick, metrics)
    elif report_type == "industry_overview":
        title = "医药生物板块财务综述"
        builder.industry_overview(conn, companies, series_map, pick, metrics)
    else:
        title = f"医药板块财报跟踪（{len(series_map)} 家公司）"
        builder.periodic_update(conn, companies, series_map, pick)

    builder.finish(title)

    report_id = repository.new_id("rpt")
    file_name = f"{report_id}.docx"
    reports_dir.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    doc.save(buffer)
    (reports_dir / file_name).write_bytes(buffer.getvalue())

    conn.execute(
        """INSERT INTO reports (id, title, report_type, codes, period_end, metrics, status, file_name, created_at)
           VALUES (?, ?, ?, ?, ?, ?, 'completed', ?, ?)""",
        (report_id, title, report_type, json.dumps(list(series_map), ensure_ascii=False), period_end,
         json.dumps(metrics, ensure_ascii=False), file_name, _now()),
    )
    report = get_report(conn, report_id)
    assert report is not None
    return report


# --------------------------------------------------------------------------- docx helpers


def _setup_document(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Microsoft YaHei")
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.8)


class _Builder:
    def __init__(self, doc: Document, source: str, settings: dict[str, str]) -> None:
        self.doc = doc
        self.source = source
        self.settings = settings
        self.title_paragraph = doc.add_paragraph()
        self.meta_paragraph = doc.add_paragraph()

    # ---- primitives
    def heading(self, text: str, level: int = 1) -> None:
        h = self.doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    def para(self, text: str, size: float = 10.5, color: tuple[int, int, int] | None = None, italic: bool = False) -> None:
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.italic = italic
        if color:
            run.font.color.rgb = RGBColor(*color)

    def table(self, header: list[str], rows: list[list[str]], note: str | None = None) -> None:
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Light Grid Accent 1"
        for i, text in enumerate(header):
            cell = table.rows[0].cells[i]
            cell.text = text
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(9)
        for row in rows:
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = text
                for p in cells[i].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                    for r in p.runs:
                        r.font.size = Pt(9)
        self.para(note or self.source, size=8, color=(0x94, 0xA3, 0xB8))

    def image(self, png: bytes) -> None:
        self.doc.add_picture(io.BytesIO(png), width=Inches(6.0))

    def finish(self, title: str) -> None:
        run = self.title_paragraph.add_run(title)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        meta = self.meta_paragraph.add_run(
            f"{self.settings['institution_name']} · {self.settings['workspace_name']} · 生成时间 {datetime.now():%Y-%m-%d %H:%M} · ResearchFlow 自动生成初稿，发布前请复核"
        )
        meta.font.size = Pt(9)
        meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        self.doc.add_paragraph()
        self.heading("口径说明", level=2)
        self.para(METHODOLOGY_NOTE, size=9, color=(0x64, 0x74, 0x8B))

    # ---- shared blocks
    def core_table(self, series: list[PeriodMetrics], n: int = 6) -> None:
        recent = series[-n:]
        header = ["指标"] + [m.quarter_label for m in recent]
        rows = []
        for key in ("revenue", "revenue_yoy", "revenue_single", "revenue_single_yoy", "net_profit", "net_profit_yoy",
                    "net_profit_single", "gross_margin", "net_margin", "selling_ratio", "rd_ratio", "operating_cash_flow"):
            metric = METRIC_MAP[key]
            flats = [flatten(m) for m in recent]
            rows.append([f"{metric.label}（{metric.unit}）"] + [format_value(f.get(key), metric, signed=metric.unit == "%") for f in flats])
        self.table(header, rows)

    def standard_charts(self, company_name: str, series: list[PeriodMetrics]) -> None:
        window = int(self.settings.get("chart_window", "8"))
        recent = series[-window:]
        labels = [m.quarter_label for m in recent]
        basis_single = self.settings.get("default_basis", "single") == "single"
        basis_label = "单季度" if basis_single else "累计"

        def pick_amount(m: PeriodMetrics, key: str) -> Number:
            return m.single.get(key) if basis_single else m.cumulative.get(key)

        def pick_growth(m: PeriodMetrics, key: str) -> Number:
            return m.yoy_single.get(key) if basis_single else m.yoy_cumulative.get(key)

        def pick_ratio(m: PeriodMetrics, key: str) -> Number:
            return m.ratios_single.get(key) if basis_single else m.ratios_cumulative.get(key)

        self.image(charts.bar_line_chart(labels, [pick_amount(m, "revenue") for m in recent], [pick_growth(m, "revenue") for m in recent],
                                         f"{company_name}：{basis_label}营业收入及同比增速", "营业收入（亿元）", "同比增速（%）", self.source))
        self.image(charts.bar_line_chart(labels, [pick_amount(m, "net_profit") for m in recent], [pick_growth(m, "net_profit") for m in recent],
                                         f"{company_name}：{basis_label}归母净利润及同比增速", "归母净利润（亿元）", "同比增速（%）", self.source))
        self.image(charts.multi_line_chart(labels, [("毛利率", [pick_ratio(m, "gross_margin") for m in recent]),
                                                    ("净利率", [pick_ratio(m, "net_margin") for m in recent]),
                                                    ("扣非净利率", [pick_ratio(m, "deducted_net_margin") for m in recent])],
                                           f"{company_name}：{basis_label}毛利率与净利率", self.source))
        self.image(charts.multi_line_chart(labels, [("销售费用率", [pick_ratio(m, "selling_ratio") for m in recent]),
                                                    ("管理费用率", [pick_ratio(m, "admin_ratio") for m in recent]),
                                                    ("研发费用率", [pick_ratio(m, "rd_ratio") for m in recent]),
                                                    ("财务费用率", [pick_ratio(m, "finance_ratio") for m in recent])],
                                           f"{company_name}：{basis_label}期间费用率", self.source))

    # ---- report types
    def deep_dive(self, company: dict, series: list[PeriodMetrics], target: PeriodMetrics, metrics: list[str]) -> None:
        self.heading("一、核心财务数据", 1)
        self.para(build_commentary(company["name"], company["code"], target))
        self.heading("二、分季度数据", 1)
        self.core_table(series)
        self.heading("三、标准图表", 1)
        self.standard_charts(company["name"], series)
        self.heading("四、历史报告期一览", 1)
        self.metric_history_table({company["code"]: company}, {company["code"]: series}, metrics, max_periods=8)

    def metric_history_table(self, companies: dict, series_map: dict[str, list[PeriodMetrics]], metrics: list[str], max_periods: int = 8) -> None:
        for key in metrics:
            metric = METRIC_MAP[key]
            all_periods = sorted({m.period_end for s in series_map.values() for m in s})[-max_periods:]
            header = ["公司"] + [Period.from_period_end(p).quarter_label for p in all_periods]
            rows = []
            for code, series in series_map.items():
                by = {m.period_end: flatten(m) for m in series}
                rows.append([companies[code]["name"]] + [format_value(by.get(p, {}).get(key), metric, signed=metric.unit == "%") for p in all_periods])
            self.para(f"{metric.label}（{metric.unit}）", size=10, color=(0x1E, 0x3A, 0x8A))
            self.table(header, rows)

    def peer_comparison(self, companies: dict, series_map: dict[str, list[PeriodMetrics]], pick, metrics: list[str]) -> None:
        latest = {code: pick(series) for code, series in series_map.items()}
        labels = [companies[c]["name"] for c in series_map]
        self.heading("一、最新报告期对比", 1)
        header = ["指标"] + [f"{companies[c]['name']}\n{latest[c].quarter_label}" for c in series_map]
        rows = []
        for key in metrics:
            metric = METRIC_MAP[key]
            rows.append([f"{metric.label}（{metric.unit}）"] + [format_value(flatten(latest[c]).get(key), metric, signed=metric.unit == "%") for c in series_map])
        self.table(header, rows)

        self.heading("二、指标排名", 1)
        for key in metrics[:4]:
            metric = METRIC_MAP[key]
            values = [(companies[c]["name"], flatten(latest[c]).get(key)) for c in series_map]
            ranked = sorted([v for v in values if v[1] is not None], key=lambda x: x[1], reverse=metric.higher_is_better)
            text = "；".join(f"{i + 1}. {name} {format_value(v, metric, signed=metric.unit == '%')}" for i, (name, v) in enumerate(ranked))
            self.para(f"{metric.label}：{text}。")

        self.heading("三、对比图表", 1)
        for key in metrics[:4]:
            metric = METRIC_MAP[key]
            self.image(charts.grouped_bar_chart(labels, [(metric.short_label, [flatten(latest[c]).get(key) for c in series_map])],
                                                f"{metric.label}（{metric.unit}）· 最新报告期", self.source, unit="%" if metric.unit == "%" else ""))
        self.heading("四、历史数据", 1)
        self.metric_history_table(companies, series_map, metrics[:4])

    def industry_overview(self, conn, companies: dict, series_map: dict[str, list[PeriodMetrics]], pick, metrics: list[str]) -> None:
        latest = {code: pick(series) for code, series in series_map.items()}
        flats = {code: flatten(m) for code, m in latest.items()}

        self.heading("一、板块整体", 1)
        def avg(key: str, codes: list[str]) -> Number:
            vals = [flats[c].get(key) for c in codes if flats[c].get(key) is not None]
            return round(sum(vals) / len(vals), 2) if vals else None
        all_codes = list(series_map)
        rev_g, np_g, gm, sr = (avg(k, all_codes) for k in ("revenue_yoy", "net_profit_yoy", "gross_margin", "selling_ratio"))
        self.para(
            f"本次综述覆盖 {len(all_codes)} 家公司。最新报告期营收同比增速平均 {format_value(rev_g, METRIC_MAP['revenue_yoy'], True)}，"
            f"归母净利润同比增速平均 {format_value(np_g, METRIC_MAP['net_profit_yoy'], True)}；平均毛利率 {format_value(gm, METRIC_MAP['gross_margin'])}，"
            f"平均销售费用率 {format_value(sr, METRIC_MAP['selling_ratio'])}。"
        )
        self.heading("二、子行业表现", 1)
        header = ["子行业", "公司数", "营收增速均值", "净利增速均值", "毛利率均值", "销售费用率均值"]
        rows = []
        for sub in SUB_INDUSTRY_ORDER:
            codes = [c for c in all_codes if classify(c) == sub]
            if not codes:
                continue
            rows.append([SUB_INDUSTRY_LABELS[sub], str(len(codes))] + [
                format_value(avg(k, codes), METRIC_MAP[k], signed=k.endswith("yoy")) for k in ("revenue_yoy", "net_profit_yoy", "gross_margin", "selling_ratio")
            ])
        self.table(header, rows)

        self.heading("三、公司一览", 1)
        header = ["公司", "报告期"] + [f"{METRIC_MAP[k].short_label}（{METRIC_MAP[k].unit}）" for k in metrics[:6]]
        rows = []
        for code in sorted(all_codes, key=lambda c: -(flats[c].get("revenue") or 0)):
            rows.append([companies[code]["name"], latest[code].quarter_label] + [format_value(flats[code].get(k), METRIC_MAP[k], signed=METRIC_MAP[k].unit == "%") for k in metrics[:6]])
        self.table(header, rows)

        self.heading("四、增速前后五名", 1)
        ranked = sorted([(companies[c]["name"], flats[c].get("revenue_yoy")) for c in all_codes if flats[c].get("revenue_yoy") is not None], key=lambda x: x[1], reverse=True)
        top = "；".join(f"{n} {format_value(v, METRIC_MAP['revenue_yoy'], True)}" for n, v in ranked[:5])
        bottom = "；".join(f"{n} {format_value(v, METRIC_MAP['revenue_yoy'], True)}" for n, v in ranked[-5:][::-1])
        self.para(f"营收增速领先：{top}。")
        self.para(f"营收增速靠后：{bottom}。")
        self.image(charts.grouped_bar_chart([n for n, _ in ranked], [("营收同比增速", [v for _, v in ranked])], "营收同比增速（%）· 最新报告期", self.source))

        self._alerts_section(conn, all_codes, "五")

    def periodic_update(self, conn, companies: dict, series_map: dict[str, list[PeriodMetrics]], pick) -> None:
        for i, (code, series) in enumerate(series_map.items(), start=1):
            target = pick(series)
            self.heading(f"{i}. {companies[code]['name']}（{code}）{target.label}", 1)
            self.para(build_commentary(companies[code]["name"], code, target))
            self.core_table(series, n=5)
        self._alerts_section(conn, list(series_map), "附")

    def _alerts_section(self, conn, codes: list[str], prefix: str) -> None:
        active = [a for a in alert_service.list_alerts(conn, status="active") if a["code"] in codes]
        self.heading(f"{prefix}、活跃告警", 1)
        if not active:
            self.para("当前没有活跃告警。")
            return
        sev = {"high": "高", "medium": "中", "low": "低"}
        self.table(["严重度", "公司", "告警", "说明"], [[sev[a["severity"]], a["company_name"] or a["code"], a["title"], a["description"]] for a in active])
