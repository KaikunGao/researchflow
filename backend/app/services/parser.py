"""Wind / Choice 导出文件解析器。

支持两种布局：
1. 宽表（Wind 数据浏览器 / 财务报表导出）：
   证券代码 | 证券简称 | 指标名称 | 2023-03-31 | 2023-06-30 | ...
2. 长表（一行一个公司-报告期）：
   证券代码 | 证券简称 | 报告期 | 营业收入 | 归母净利润 | ...

输出统一为 ParsedFact（金额已换算为亿元）。
"""

import csv
import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.services.fields import (
    extract_unit,
    extract_unit_declaration,
    match_field,
    to_yi,
)
from app.services.periods import parse_period

_CODE_RE = re.compile(r"^(?:(SH|SZ|BJ)\.?)?(\d{6})(?:\.(SH|SZ|BJ))?$", re.IGNORECASE)
_PERIOD_HEADER_RE = re.compile(r"报告期|报告日期|截止日期|日期|期末|报表日期|period|date", re.IGNORECASE)
_NAME_HEADER_RE = re.compile(r"简称|名称|公司|name", re.IGNORECASE)
_INDICATOR_HEADER_RE = re.compile(r"指标|项目|科目|报表项目|item|indicator", re.IGNORECASE)
_CODE_HEADER_RE = re.compile(r"代码|code", re.IGNORECASE)


@dataclass
class ParsedFact:
    code: str
    name: str
    period_end: str
    field: str
    value: float


@dataclass
class ParseResult:
    facts: list[ParsedFact] = field(default_factory=list)
    layout: str = "unknown"
    row_count: int = 0
    warnings: list[str] = field(default_factory=list)
    unmapped_indicators: list[str] = field(default_factory=list)

    @property
    def company_count(self) -> int:
        return len({f.code for f in self.facts})

    @property
    def period_count(self) -> int:
        return len({f.period_end for f in self.facts})


class ParseError(ValueError):
    """文件无法解析时抛出，message 面向用户。"""


# --------------------------------------------------------------------------- #
#                                  入口                                        #
# --------------------------------------------------------------------------- #


def parse_file(file_name: str, content: bytes) -> ParseResult:
    lower = file_name.lower()
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        sheets = _read_xlsx(content)
    elif lower.endswith(".csv"):
        sheets = [("csv", _read_csv(content))]
    elif lower.endswith(".xls"):
        raise ParseError("暂不支持旧版 .xls 文件，请在 Excel 中另存为 .xlsx 后再导入。")
    else:
        raise ParseError("仅支持 .xlsx / .csv 文件。")

    result = ParseResult()
    parsed_any = False
    for sheet_name, rows in sheets:
        if not rows:
            continue
        sheet_result = _parse_sheet(sheet_name, rows)
        if sheet_result is None:
            result.warnings.append(f"工作表「{sheet_name}」未识别出财务数据布局，已跳过。")
            continue
        parsed_any = True
        result.facts.extend(sheet_result.facts)
        result.row_count += sheet_result.row_count
        result.warnings.extend(sheet_result.warnings)
        for name in sheet_result.unmapped_indicators:
            if name not in result.unmapped_indicators:
                result.unmapped_indicators.append(name)
        if result.layout == "unknown":
            result.layout = sheet_result.layout
        elif result.layout != sheet_result.layout:
            result.layout = "mixed"

    if not parsed_any:
        raise ParseError(
            "未识别出可解析的财务数据。请确认文件包含证券代码列，以及报告期（如 2024-12-31）作为列头或「报告期」列。"
        )
    if not result.facts:
        raise ParseError("文件已读取，但没有解析出任何可识别的财务指标。请检查指标名称是否为营业收入、归母净利润等常见科目。")

    # 去重：同一 (code, period, field) 以最后出现的为准
    dedup: dict[tuple[str, str, str], ParsedFact] = {}
    for fact in result.facts:
        dedup[(fact.code, fact.period_end, fact.field)] = fact
    result.facts = list(dedup.values())
    if result.unmapped_indicators:
        shown = "、".join(result.unmapped_indicators[:8])
        more = f" 等 {len(result.unmapped_indicators)} 项" if len(result.unmapped_indicators) > 8 else ""
        result.warnings.append(f"以下指标未能映射到标准字段，已忽略：{shown}{more}。")
    return result


# --------------------------------------------------------------------------- #
#                                  文件读取                                    #
# --------------------------------------------------------------------------- #


def _read_xlsx(content: bytes) -> list[tuple[str, list[list[object]]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"无法读取 Excel 文件：{exc}") from exc
    sheets: list[tuple[str, list[list[object]]]] = []
    for ws in workbook.worksheets:
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        sheets.append((ws.title, rows))
    workbook.close()
    return sheets


def _read_csv(content: bytes) -> list[list[object]]:
    text: str | None = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ParseError("无法识别 CSV 文件编码，请另存为 UTF-8 后重试。")
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [list(row) for row in reader]


# --------------------------------------------------------------------------- #
#                                  布局识别                                    #
# --------------------------------------------------------------------------- #


def _parse_sheet(sheet_name: str, rows: list[list[object]]) -> ParseResult | None:
    rows = [[_clean_cell(c) for c in row] for row in rows]
    rows = [row for row in rows if any(c not in (None, "") for c in row)]
    if not rows:
        return None

    default_unit = _detect_unit_declaration(rows[:10])
    scan_limit = min(len(rows), 20)

    for idx in range(scan_limit):
        header = rows[idx]
        period_cols = {i: p for i, c in enumerate(header) if (p := parse_period(c)) is not None}
        if len(period_cols) >= 2:
            return _parse_wide(sheet_name, rows, idx, period_cols, default_unit)

    for idx in range(scan_limit):
        header = rows[idx]
        period_col = _find_col(header, _PERIOD_HEADER_RE)
        field_cols = {i: key for i, c in enumerate(header) if isinstance(c, str) and (key := match_field(c))}
        if period_col is not None and len(field_cols) >= 1:
            return _parse_long(sheet_name, rows, idx, period_col, field_cols, default_unit)

    return None


def _parse_wide(
    sheet_name: str,
    rows: list[list[object]],
    header_idx: int,
    period_cols: dict,
    default_unit: str | None,
) -> ParseResult:
    header = rows[header_idx]
    body = rows[header_idx + 1 :]
    non_period = [i for i in range(len(header)) if i not in period_cols]

    code_col = _find_code_col(body, non_period)
    name_col = _find_col(header, _NAME_HEADER_RE, candidates=non_period)
    indicator_col = _find_col(header, _INDICATOR_HEADER_RE, candidates=non_period)
    if indicator_col is None:
        # 取可映射指标最多的那一列
        best, best_hits = None, 0
        for i in non_period:
            if i in (code_col, name_col):
                continue
            hits = sum(1 for row in body if i < len(row) and isinstance(row[i], str) and match_field(row[i]))
            if hits > best_hits:
                best, best_hits = i, hits
        indicator_col = best
    if indicator_col is None:
        return None

    sheet_code, sheet_name_guess = (None, None)
    if code_col is None:
        sheet_code, sheet_name_guess = _detect_sheet_company(sheet_name, rows[:header_idx])
        if sheet_code is None:
            result = ParseResult(layout="wide")
            result.warnings.append(f"工作表「{sheet_name}」缺少证券代码列，已跳过。")
            return result

    result = ParseResult(layout="wide")
    unit_default = default_unit or "元"
    unit_warned = False

    for row in body:
        indicator_raw = row[indicator_col] if indicator_col < len(row) else None
        if not isinstance(indicator_raw, str):
            continue
        key = match_field(indicator_raw)
        if key is None:
            if indicator_raw.strip() and indicator_raw not in result.unmapped_indicators:
                result.unmapped_indicators.append(indicator_raw.strip())
            continue
        code = _normalize_code(row[code_col]) if code_col is not None and code_col < len(row) else sheet_code
        if code is None:
            continue
        name = _cell_str(row[name_col]) if name_col is not None and name_col < len(row) else (sheet_name_guess or code)
        unit = extract_unit(indicator_raw) or default_unit
        if unit is None and not unit_warned:
            result.warnings.append(f"工作表「{sheet_name}」未声明金额单位，已默认按「元」换算为亿元。")
            unit_warned = True
        result.row_count += 1
        for col, period in period_cols.items():
            value = _to_number(row[col]) if col < len(row) else None
            if value is None:
                continue
            result.facts.append(ParsedFact(code, name or code, period.period_end, key, round(to_yi(value, unit, unit_default), 6)))
    return result


def _parse_long(
    sheet_name: str,
    rows: list[list[object]],
    header_idx: int,
    period_col: int,
    field_cols: dict[int, str],
    default_unit: str | None,
) -> ParseResult:
    header = rows[header_idx]
    body = rows[header_idx + 1 :]
    others = [i for i in range(len(header)) if i != period_col and i not in field_cols]
    code_col = _find_code_col(body, others)
    name_col = _find_col(header, _NAME_HEADER_RE, candidates=others)

    sheet_code, sheet_name_guess = (None, None)
    if code_col is None:
        sheet_code, sheet_name_guess = _detect_sheet_company(sheet_name, rows[:header_idx])
        if sheet_code is None:
            result = ParseResult(layout="long")
            result.warnings.append(f"工作表「{sheet_name}」缺少证券代码列，已跳过。")
            return result

    result = ParseResult(layout="long")
    unit_default = default_unit or "元"
    col_units = {i: (extract_unit(str(header[i])) or default_unit) for i in field_cols}
    if any(u is None for u in col_units.values()):
        result.warnings.append(f"工作表「{sheet_name}」部分指标列未声明金额单位，已默认按「元」换算为亿元。")
    for i, c in enumerate(header):
        if i in field_cols or i == period_col or i in (code_col, name_col):
            continue
        if isinstance(c, str) and c.strip() and not _CODE_HEADER_RE.search(c):
            result.unmapped_indicators.append(c.strip())

    for row in body:
        period = parse_period(row[period_col]) if period_col < len(row) else None
        if period is None:
            continue
        code = _normalize_code(row[code_col]) if code_col is not None and code_col < len(row) else sheet_code
        if code is None:
            continue
        name = _cell_str(row[name_col]) if name_col is not None and name_col < len(row) else (sheet_name_guess or code)
        result.row_count += 1
        for col, key in field_cols.items():
            value = _to_number(row[col]) if col < len(row) else None
            if value is None:
                continue
            result.facts.append(
                ParsedFact(code, name or code, period.period_end, key, round(to_yi(value, col_units[col], unit_default), 6))
            )
    return result


# --------------------------------------------------------------------------- #
#                                  工具函数                                    #
# --------------------------------------------------------------------------- #


def _clean_cell(value: object) -> object:
    if isinstance(value, str):
        text = value.strip()
        return text if text else None
    return value


def _cell_str(value: object) -> str | None:
    if value is None:
        return None
    return str(value).strip()


def _to_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    if text in ("", "--", "-", "—", "N/A", "NA", "null", "None"):
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


def _normalize_code(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip().upper()
    if isinstance(value, float) and value.is_integer():
        text = f"{int(value):06d}"
    match = _CODE_RE.match(text)
    if not match:
        return None
    prefix, digits, suffix = match.groups()
    exchange = (suffix or prefix or _infer_exchange(digits)).upper()
    return f"{digits}.{exchange}"


def _infer_exchange(digits: str) -> str:
    if digits.startswith(("6", "9")):
        return "SH"
    if digits.startswith(("0", "2", "3")):
        return "SZ"
    return "BJ"


def _find_col(header: list[object], pattern: re.Pattern[str], candidates: list[int] | None = None) -> int | None:
    indices = candidates if candidates is not None else range(len(header))
    for i in indices:
        cell = header[i] if i < len(header) else None
        if isinstance(cell, str) and pattern.search(cell):
            return i
    return None


def _find_code_col(body: list[list[object]], candidates: list[int]) -> int | None:
    best, best_hits = None, 0
    for i in candidates:
        hits = sum(1 for row in body[:50] if i < len(row) and _normalize_code(row[i]) is not None)
        if hits > best_hits:
            best, best_hits = i, hits
    return best if best_hits > 0 else None


def _detect_unit_declaration(rows: list[list[object]]) -> str | None:
    for row in rows:
        for cell in row:
            if isinstance(cell, str):
                unit = extract_unit_declaration(cell)
                if unit:
                    return unit
    return None


def _detect_sheet_company(sheet_name: str, title_rows: list[list[object]]) -> tuple[str | None, str | None]:
    """单公司工作表：从 sheet 名或标题行里找证券代码，例如 '恒瑞医药(600276.SH)'。"""
    texts = [sheet_name] + [str(c) for row in title_rows for c in row if isinstance(c, str)]
    for text in texts:
        for token in re.findall(r"(?:SH|SZ|BJ)?\.?\d{6}(?:\.(?:SH|SZ|BJ))?", text, flags=re.IGNORECASE):
            code = _normalize_code(token)
            if code:
                before, _, after = text.partition(token)
                name = re.sub(r"[（()）\s]+$", "", before).strip() or re.sub(r"^[（()）\s]+", "", after).strip()
                name = name.split()[0] if name else ""
                return code, (name or None)
    return None, None
