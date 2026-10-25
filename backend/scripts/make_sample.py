"""生成 Wind 风格的示例导出文件（演示数据）。

运行：python scripts/make_sample.py
输出：samples/wind_sample_pharma.xlsx

所有数值由确定性公式生成，仅用于演示导入 → 拆单季 → 速览的完整流程，
不代表任何公司的真实财务状况。
"""

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "samples" / "wind_sample_pharma.xlsx"

# 公司画像：2021Q4 单季营收基数（亿元）、年化增速、季节性、成本率、费用率、净利率、资产规模
COMPANIES = [
    ("600276.SH", "恒瑞医药", 62.0, 0.12, [0.95, 1.00, 1.02, 1.03], 0.14, 0.33, 0.04, 0.24, 0.20, 0.13, 420.0, 0.12),
    ("300760.SZ", "迈瑞医疗", 90.0, 0.14, [0.90, 1.05, 1.00, 1.05], 0.36, 0.15, 0.04, 0.10, 0.32, 0.36, 480.0, 0.34),
    ("603259.SH", "药明康德", 100.0, 0.08, [0.88, 1.02, 1.04, 1.06], 0.59, 0.03, 0.07, 0.04, 0.23, 0.30, 720.0, 0.28),
    ("600436.SH", "片仔癀", 24.0, 0.10, [1.05, 1.00, 0.98, 0.97], 0.54, 0.09, 0.03, 0.03, 0.28, 0.17, 130.0, 0.18),
    ("300015.SZ", "爱尔眼科", 42.0, 0.13, [0.92, 1.02, 1.10, 0.96], 0.49, 0.10, 0.13, 0.02, 0.17, 0.30, 280.0, 0.31),
    ("000661.SZ", "长春高新", 32.0, 0.06, [0.95, 1.05, 1.02, 0.98], 0.09, 0.30, 0.06, 0.14, 0.33, 0.21, 250.0, 0.23),
]

PERIODS = [(y, q) for y in (2022, 2023, 2024) for q in (1, 2, 3, 4)]

# 2024 年业绩压力（营收乘数、净利率变化），用于演示告警：
# 长春高新：生长激素集采 → 营收下滑、净利率大幅收窄；药明康德：海外需求放缓 → 营收小幅下滑
STRESS = {
    "000661.SZ": {(2024, 1): (0.95, -0.03), (2024, 2): (0.88, -0.06), (2024, 3): (0.82, -0.09), (2024, 4): (0.78, -0.12)},
    "603259.SH": {(2024, 1): (0.96, -0.01), (2024, 2): (0.94, -0.01), (2024, 3): (0.92, -0.02), (2024, 4): (0.90, -0.02)},
}
QUARTER_END = {1: "03-31", 2: "06-30", 3: "09-30", 4: "12-31"}

INDICATORS = [
    ("营业收入", "revenue"),
    ("营业成本", "operating_cost"),
    ("归属母公司股东的净利润", "net_profit"),
    ("扣除非经常性损益后的净利润", "net_profit_deducted"),
    ("销售费用", "selling_expense"),
    ("管理费用", "admin_expense"),
    ("研发费用", "rd_expense"),
    ("财务费用", "finance_expense"),
    ("经营活动产生的现金流量净额", "operating_cash_flow"),
    ("资产总计", "total_assets"),
    ("负债合计", "total_liabilities"),
    ("归属母公司股东的权益", "equity_parent"),
    ("存货", "inventory"),
    ("应收账款", "accounts_receivable"),
    ("每股收益(摊薄)", "unmapped"),  # 故意保留一个无法映射的指标，演示告警
]


def quarterly_values(profile):
    code, name, base, growth, season, cost_r, sell_r, adm_r, rd_r, nm, fin_r_dummy, assets0, debt_r = profile
    del fin_r_dummy, name
    out = []
    for i, (y, q) in enumerate(PERIODS):
        t = i + 1
        shock = 1 + 0.06 * math.sin(t * 1.3) - (0.08 if (y == 2024 and q >= 3 and growth < 0.1) else 0)
        rev = base * ((1 + growth) ** (t / 4)) * season[q - 1] * shock
        rev_mult, margin_delta = STRESS.get(profile[0], {}).get((y, q), (1.0, 0.0))
        rev *= rev_mult
        cost = rev * (cost_r + 0.01 * math.sin(t))
        selling = rev * (sell_r + 0.004 * math.cos(t))
        admin = rev * adm_r
        rd = rev * (rd_r + 0.003 * t / 12)
        finance = rev * (-0.004 if debt_r < 0.25 else 0.006)
        net_margin = nm - 0.02 * math.sin(t * 0.9) - (0.03 if y == 2024 and q == 4 else 0) + margin_delta
        net_profit = rev * net_margin
        deducted = net_profit * 0.93
        ocf = net_profit * (1.1 + 0.3 * math.sin(t * 0.7))
        assets = assets0 * (1 + 0.025) ** t
        liabilities = assets * (debt_r + 0.01 * math.sin(t * 0.5))
        equity = (assets - liabilities) * 0.96
        inventory = rev * 4 * (0.10 if cost_r < 0.3 else 0.18)
        receivable = rev * 4 * 0.12
        out.append(
            {
                "revenue": rev,
                "operating_cost": cost,
                "net_profit": net_profit,
                "net_profit_deducted": deducted,
                "selling_expense": selling,
                "admin_expense": admin,
                "rd_expense": rd,
                "finance_expense": finance,
                "operating_cash_flow": ocf,
                "total_assets": assets,
                "total_liabilities": liabilities,
                "equity_parent": equity,
                "inventory": inventory,
                "accounts_receivable": receivable,
                "unmapped": net_profit / 6,
            }
        )
    return out


FLOW = {"revenue", "operating_cost", "net_profit", "net_profit_deducted", "selling_expense", "admin_expense",
        "rd_expense", "finance_expense", "operating_cash_flow"}


def cumulative(values):
    out = []
    for i, (y, q) in enumerate(PERIODS):
        row = {}
        for key, v in values[i].items():
            if key in FLOW:
                row[key] = sum(values[j][key] for j in range(i - q + 1, i + 1))
            else:
                row[key] = v
        out.append(row)
    return out


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "财务数据"
    ws["A1"] = "数据来源：Wind（演示数据，非真实财务数据）   单位：元"
    ws["A1"].font = Font(color="666666")
    header = ["证券代码", "证券简称", "指标名称"] + [f"{y}-{QUARTER_END[q]}" for y, q in PERIODS]
    ws.append([])
    ws.append(header)
    for c in ws[3]:
        c.font = Font(bold=True)

    for profile in COMPANIES:
        code, name = profile[0], profile[1]
        cum = cumulative(quarterly_values(profile))
        for label, key in INDICATORS:
            row = [code, name, label] + [round(c[key] * 1e8, 2) for c in cum]
            ws.append(row)

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 28

    note = wb.create_sheet("说明")
    note["A1"] = "本文件为 ResearchFlow 示例文件，模拟 Wind 财务报表导出的宽表格式。"
    note["A2"] = "所有财务数值由确定性公式生成，仅用于演示导入、单季度拆分与财报速览流程，不代表任何公司的真实财务状况。"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
